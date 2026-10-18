"""S8-09 (texts and file, part B): what the daily monitoring writes, and the list as a file.

Comparisons are made by ``compare`` itself from snapshots the adapters build, as in the
rules' own tests; names and INNs are synthetic.
"""

import io
from dataclasses import replace
from datetime import UTC, date, datetime

import pytest
from openpyxl import load_workbook

from claims_assistant.domain.external import Section
from claims_assistant.domain.monitoring import compare
from claims_assistant.domain.watchlist import WatchEntry
from claims_assistant.infrastructure.excel.watchlist import SHEET_NAME, TITLES, build_watchlist
from claims_assistant.presentation.telegram.monitoring import (
    MAX_CHANGES,
    MESSAGE_LIMIT,
    Checked,
    format_changes,
    format_daily_summary,
)
from claims_assistant.presentation.telegram.texts import watch_frequency_prompt
from tests.unit.test_monitoring_changes import (
    INN,
    TODAY,
    YESTERDAY,
    case,
    court_snapshot,
    failed,
    full_day,
)

DAY = TODAY.date()
NAME = "ООО «Ромашка»"


def with_new_case(claim=500_000):
    old = [case("А00-1/2026", "2026-09-01")]
    new = [case("А00-2/2026", "2026-10-15", claim=claim), *old]
    return compare(
        INN,
        full_day(YESTERDAY, courts=court_snapshot(old, YESTERDAY)),
        full_day(TODAY, courts=court_snapshot(new, TODAY)),
    )


def quiet():
    return compare(INN, full_day(YESTERDAY), full_day(TODAY))


def unread():
    return compare(
        INN, full_day(YESTERDAY), full_day(TODAY, enforcements=failed(Section.ENFORCEMENTS, TODAY))
    )


def first():
    return compare(INN, [], full_day(TODAY))


# --- one company ---------------------------------------------------------------------


def test_a_change_names_the_company_the_record_its_date_and_link():
    text = format_changes(INN, NAME, with_new_case(), DAY)
    lines = text.splitlines()
    assert lines[0] == f"🔔 Мониторинг, 16.10.2026: {NAME}, ИНН {INN}"
    assert lines[2].startswith("• Новое арбитражное дело, организация — ответчик: А00-2/2026")
    assert lines[2].endswith("(15.10.2026)")
    assert lines[3] == "  https://kad.arbitr.ru/Card/uuid-А00-2/2026"


def test_nothing_worth_writing_is_no_message():
    assert format_changes(INN, NAME, quiet(), DAY) is None
    assert format_changes(INN, NAME, with_new_case(claim=99_999), DAY) is None  # below 100 000


def test_a_section_not_read_today_is_said_with_its_reason():
    old = [case("А00-1/2026", "2026-09-01")]
    new = [case("А00-2/2026", "2026-10-15"), *old]
    comparison = compare(
        INN,
        full_day(YESTERDAY, courts=court_snapshot(old, YESTERDAY)),
        full_day(
            TODAY,
            courts=court_snapshot(new, TODAY),
            enforcements=failed(Section.ENFORCEMENTS, TODAY),
        ),
    )
    text = format_changes(INN, None, comparison, DAY)
    assert text.splitlines()[0].endswith(f": ИНН {INN}")
    assert "Не проверено сегодня, изменения там неизвестны:\n  ФССП — источник не ответил" in text


def test_a_long_list_is_cut_and_the_message_fits_telegram():
    old = []
    new = [case(f"А00-{n}/2026", "2026-10-15", claim=200_000) for n in range(60)]
    comparison = compare(
        INN,
        full_day(YESTERDAY, courts=court_snapshot(old, YESTERDAY)),
        full_day(TODAY, courts=court_snapshot(new, TODAY)),
    )
    text = format_changes(INN, NAME, comparison, DAY)
    assert f"• …и ещё {60 - MAX_CHANGES}" in text
    assert len(text) <= MESSAGE_LIMIT


# --- the day -------------------------------------------------------------------------


def test_a_quiet_day_says_so_only_when_everything_was_read():
    text = format_daily_summary(DAY, [Checked(INN, NAME, quiet())])
    assert text.splitlines()[0] == "📋 Мониторинг на 16.10.2026. Организаций проверено: 1."
    assert "Изменений нет." in text


def test_a_failed_source_is_never_a_quiet_day():
    text = format_daily_summary(DAY, [Checked(INN, NAME, quiet()), Checked(INN, None, unread())])
    assert "Изменений нет." not in text
    assert "Новых изменений не найдено, но проверено не всё" in text
    assert "Проверено не полностью: 1 организация.\n  ФССП — источник не ответил: 1" in text


def test_the_day_names_who_changed_and_who_was_checked_first():
    text = format_daily_summary(
        DAY,
        [
            Checked(INN, NAME, with_new_case()),
            Checked("7707083893", "ПАО «Новое»", first()),
            Checked(INN, None, quiet()),
        ],
    )
    assert f"Изменения у 1 организации — подробности выше:\n  {NAME}, ИНН {INN}" in text
    assert "Впервые проверено: 1 организация. Это точка отсчёта" in text
    assert "Изменений нет." not in text


def rebased():
    """Yesterday Fedresurs did not answer; today it did — compared everywhere else."""
    return compare(
        INN,
        full_day(YESTERDAY, fedresurs=failed(Section.FEDRESURS, YESTERDAY)),
        full_day(TODAY),
    )


def test_a_section_that_failed_yesterday_is_not_a_company_checked_first():
    """Live run of #105 (review A on #101): the company was checked the day before."""
    text = format_daily_summary(DAY, [Checked(INN, NAME, rebased())])
    assert "Впервые проверено" not in text
    assert "Точка отсчёта по разделам, которые в прошлый раз не ответили:\n  Федресурс: 1" in text
    # Everything answered today: not «проверено не всё», and not «Изменений нет» either.
    assert "Новых изменений не найдено." in text
    assert "проверено не всё" not in text and "Изменений нет." not in text


def test_checko_keeps_its_capital_in_a_reason():
    snapshot = failed(Section.ENFORCEMENTS, TODAY)
    reason = "Checko не ответил за отведённое время."
    snapshot = replace(snapshot, missing=(reason,), error=replace(snapshot.error, message=reason))
    comparison = compare(INN, full_day(YESTERDAY), full_day(TODAY, enforcements=snapshot))
    text = format_daily_summary(DAY, [Checked(INN, NAME, comparison)])
    assert "ФССП — Checko не ответил за отведённое время: 1" in text


def test_a_date_already_in_the_label_is_not_repeated():
    comparison = with_new_case()
    change = comparison.to_notify[0]
    dated = replace(change, text=f"{change.text} от 15.10.2026")
    text = format_changes(INN, NAME, replace(comparison, changes=(dated,)), DAY)
    assert text.count("15.10.2026") == 1


def test_an_empty_day_says_the_list_is_empty():
    assert format_daily_summary(DAY, []) == "📋 Мониторинг на 16.10.2026: в списке нет организаций."


# --- the file ------------------------------------------------------------------------


def test_the_list_as_a_file_newest_first_with_text_inns_and_dates():
    entries = [
        WatchEntry("0123456789", date(2026, 10, 1), "ООО «Старое»", date(2026, 10, 16)),
        WatchEntry(
            "500100732259",
            date(2026, 10, 16),
            "ИП Иванов",
            date(2026, 10, 16),
            "Новое исполнительное производство (ФССП): № 1/26/00001-ИП",
            date(2026, 10, 15),
        ),
    ]
    sheet = load_workbook(io.BytesIO(build_watchlist(entries, DAY)))[SHEET_NAME]
    assert sheet["A1"].value == "Список на мониторинге на 16.10.2026: 2 из 100"
    assert tuple(cell.value for cell in sheet[2]) == TITLES
    newest, oldest = sheet[3], sheet[4]
    assert newest[1].value == "500100732259" and newest[5].value.startswith("Новое исполнительное")
    assert oldest[1].value == "0123456789"  # the leading zero survives
    assert oldest[3].value.date() == date(2026, 10, 1)
    assert oldest[3].number_format == "DD.MM.YYYY"
    assert sheet.freeze_panes == "C3"


def test_one_day_keeps_the_order_of_adding_not_of_the_inn():
    """The store gives the list newest first; within a day the file keeps that order."""
    entries = [
        WatchEntry("9999999999", DAY, "Добавлена последней"),
        WatchEntry("0000000001", DAY, "Добавлена раньше"),
    ]
    sheet = load_workbook(io.BytesIO(build_watchlist(entries, DAY)))[SHEET_NAME]
    assert [sheet[3][1].value, sheet[4][1].value] == ["9999999999", "0000000001"]


def test_an_empty_list_is_a_file_with_its_header():
    sheet = load_workbook(io.BytesIO(build_watchlist([], DAY)))[SHEET_NAME]
    assert sheet["A1"].value.endswith(": 0 из 100")
    assert sheet.max_row == 2


# --- S8-12: a gap after a long break, the price of each frequency --------------------


def gapped():
    week_ago = datetime(2026, 10, 9, 9, tzinfo=UTC)
    old = [case("А00-1/2026", "2026-10-01")]
    new = [case("А00-2/2026", "2026-10-12")]
    return compare(
        INN,
        full_day(week_ago, courts=court_snapshot(old, week_ago, total=150)),
        full_day(TODAY, courts=court_snapshot(new, TODAY, total=400)),
    )


def test_a_gap_is_said_under_the_changes():
    text = format_changes(INN, NAME, gapped(), DAY)
    assert "Сравнено не полностью:" in text
    assert (
        "  арбитражные дела — источник отдаёт только новейшие записи, а с прошлой проверки "
        "(09.10.2026) их прибавилось больше: появившиеся до 12.10.2026 могли не попасть в "
        "список"
    ) in text.splitlines()


def test_a_gap_is_never_a_quiet_day():
    summary = format_daily_summary(
        DAY,
        [Checked(INN, NAME, quiet()), Checked("7707083893", None, replace(gapped(), changes=()))],
    )
    assert "Изменений нет." not in summary
    assert "Новых изменений не найдено, но проверено не всё — ниже." in summary
    assert "Сравнено не полностью: 1 организация." in summary
    assert "  ИНН 7707083893: арбитражные дела" in summary.splitlines()


@pytest.mark.parametrize(
    ("count", "lines"),
    [
        (
            100,
            [
                "Для 100 организаций в списке:",
                "• каждый день: ≈ 500 запросов в день, около 1 800 ₽ в месяц;",
                "• по будням: ≈ 500 запросов в будний день, около 1 290 ₽ в месяц;",
                "• раз в неделю: ≈ 71 запрос в день, бесплатно.",
            ],
        ),
        (14, ["• каждый день: ≈ 70 запросов в день, бесплатно;"]),
    ],
)
def test_the_frequency_question_names_the_price_of_each_answer(count, lines):
    text = watch_frequency_prompt(count).splitlines()
    assert all(line in text for line in lines)


def test_an_empty_list_has_no_price():
    assert "₽" not in watch_frequency_prompt(0)
