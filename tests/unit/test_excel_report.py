"""S2-05: the four-sheet Excel report — readable, ordered, source-linked, formula-safe."""

import io
from datetime import UTC, date, datetime
from decimal import Decimal
from hashlib import sha256

import pytest
from openpyxl import load_workbook

from claims_assistant.domain.analysis import FileKind, UploadedFile
from claims_assistant.domain.counterparties import CounterpartyRow
from claims_assistant.domain.external import DataMode, Period
from claims_assistant.domain.imports import ImportIssue, IssueSeverity
from claims_assistant.domain.report import (
    DEMO_SCORE_NOTE,
    AnalysisReport,
    ReportMeta,
    ReportRow,
    demo_assessment,
    file_labels,
)
from claims_assistant.domain.scoring import Priority, assess
from claims_assistant.infrastructure.checko import bankruptcy, finances
from claims_assistant.infrastructure.checko.company_data import normalize_company
from claims_assistant.infrastructure.excel.control_package import DEMO_SCORES, demo_report
from claims_assistant.infrastructure.excel.report import PRIORITY_LABELS, SHEETS, build_report

NOW = datetime(2026, 9, 19, 9, 30, tzinfo=UTC)
CHECKED = datetime(2026, 9, 18, 12, 0, tzinfo=UTC)
DAY = date(2026, 9, 1)
INN_A, INN_B = "7700000009", "1234567894"


def company(inn, name="ДЕМО"):
    payload = {
        "meta": {"status": "ok"},
        "data": {
            "ИНН": inn,
            "ОГРН": "0000000000000",
            "НаимСокр": name,
            "ДатаВып": "2026-09-01",
            "Статус": {"Код": "001", "Наим": "Действует"},
        },
    }
    return normalize_company(payload, inn, CHECKED)


def live_row(inn, debt, overdue, records=(), years=None, name=None):
    line = CounterpartyRow(
        inn=inn,
        name=name,
        cutoff_date=DAY,
        debt=debt,
        overdue_days=overdue,
        last_payment_date=date(2026, 8, 20),
    )
    snapshots = (
        company(inn),
        bankruptcy.project_bankruptcy(inn, list(records), CHECKED, complete=True, unreadable=0),
        finances.project_finances(
            inn,
            years or {2024: {"2110": 1000, "2400": 10}, 2025: {"2110": 1100, "2400": 20}},
            CHECKED,
        ),
    )
    return ReportRow(line, assess(inn, snapshots, line), snapshots)


def live_report(rows, issues=()):
    meta = ReportMeta(
        run_id="run-1",
        analysis_date=DAY,
        mode=DataMode.LIVE,
        created_at=NOW,
        checked_at=CHECKED,
        package=("Контрагенты: 2 строки",),
    )
    return AnalysisReport(meta, tuple(rows), tuple(issues))


def open_report(report):
    return load_workbook(io.BytesIO(build_report(report)), data_only=False)


def data_rows(sheet):
    return [row for row in sheet.iter_rows(min_row=3, values_only=True)]


def test_file_reads_back_with_four_sheets_and_a_fixed_layout():
    workbook = open_report(live_report([live_row(INN_B, Decimal("100.00"), 5)]))
    assert workbook.sheetnames == list(SHEETS)
    priorities = workbook["Приоритеты"]
    assert priorities["A1"].value.startswith("Приоритеты — проверка run-1, дата анализа 01.09.2026")
    assert [c.value for c in priorities[2]] == [
        "ИНН",
        "Название",
        "Долг, ₽",
        "Дней просрочки",
        "Приоритет",
        "Полнота",
        "Причины",
        "Следующий шаг",
        "Пояснение",
        "Обещания оплаты",
        "Договоров",
        "Худшая просрочка по договору, дн.",
    ]
    assert priorities.freeze_panes == "A3"


def test_priorities_follow_the_report_order_and_explain_the_signal():
    event = [{"Дата": "2026-06-01", "Тип": "Введение наблюдения", "Номер": "E-1"}]
    rows = [
        live_row(INN_A, Decimal("900.00"), 0),  # clean → low
        live_row(INN_B, Decimal("10.00"), 5, records=event),  # EFRSB message → high
    ]
    sheet = open_report(live_report(rows))["Приоритеты"]
    (first, second) = data_rows(sheet)
    assert first[0] == INN_B and first[4] == PRIORITY_LABELS[Priority.HIGH]
    assert "ЕФРСБ" in first[6] and "efrsb-event-0" in first[6]
    assert "Проверить событие" in first[7]
    assert second[0] == INN_A and second[4] == PRIORITY_LABELS[Priority.LOW]
    assert second[5] == "полная"
    assert second[2] == 900  # stored as a number, not text


def test_grounds_list_internal_and_external_facts_with_sources():
    sheet = open_report(live_report([live_row(INN_B, Decimal("100.00"), 5)]))["Основания"]
    rows = data_rows(sheet)
    ids = {row[0] for row in rows}
    assert {"internal-debt", "internal-overdue", "company-name", "revenue-2025"} <= ids
    revenue = next(row for row in rows if row[0] == "revenue-2025")
    assert revenue[1] == INN_B
    assert revenue[3] == "1 100,00 ₽"
    assert revenue[4] == "01.01.2025–31.12.2025"
    assert revenue[5] == "checko-finances-v2"
    debt = next(row for row in rows if row[0] == "internal-debt")
    assert debt[5] == "Файл «Контрагенты»"


def test_quality_lists_import_issues_without_inn_and_data_gaps_per_inn():
    issue = ImportIssue(
        code="inn_invalid",
        severity=IssueSeverity.ERROR,
        sheet="Контрагенты",
        reason="Контрольная цифра ИНН не совпадает; проверьте значение.",
        row=7,
        column="A",
    )
    unknown = live_row(INN_B, None, None, years={})  # gaps, no signals
    rows = data_rows(open_report(live_report([unknown], [issue]))["Качество данных"])
    assert rows[0] == (
        "Ошибка строки",
        None,
        "лист «Контрагенты», строка 7, колонка A",
        issue.reason,
    )
    gaps = [row for row in rows if row[0] == "Пропуск или ограничение данных"]
    assert gaps and all(row[1] == INN_B for row in gaps)


def test_file_labels_name_kind_period_and_number_of_repeated_kinds():
    def upload(file_id, kind, coverage=None):
        return UploadedFile(
            id=file_id,
            run_id="run-1",
            kind=kind,
            checksum=sha256(file_id.encode()).hexdigest(),
            size_bytes=1,
            stored_path=f"run-1/{file_id}.xlsx",
            uploaded_at=NOW,
            coverage=coverage,
        )

    june = Period(date(2026, 6, 1), date(2026, 6, 30))
    files = [
        upload("a", FileKind.COUNTERPARTIES),
        upload("b", FileKind.DEBT_HISTORY),
        upload("c", FileKind.PAYMENTS, june),
        upload("d", FileKind.DEBT_HISTORY),
    ]
    assert file_labels(files) == (
        ("a", "Контрагенты"),
        ("b", "История долга, файл 1"),
        ("c", "Платежи за 01.06.2026–30.06.2026"),
        ("d", "История долга, файл 2"),
    )


def test_quality_names_the_file_of_each_issue_in_upload_order():
    files = (("f-main", "Контрагенты"), ("f-pay-1", "Платежи за 01.06.2026–30.06.2026, файл 1"))
    files += (("f-pay-2", "Платежи за 01.07.2026–31.08.2026, файл 2"),)
    meta = ReportMeta("run-1", DAY, DataMode.LIVE, NOW, CHECKED, files=files)

    def issue(file_id, row, sheet="Платежи"):
        return ImportIssue(
            code="x",
            severity=IssueSeverity.WARNING,
            sheet=sheet,
            reason="Причина.",
            file_id=file_id,
            row=row,
        )

    issues = (
        issue("f-pay-2", 3),
        issue("f-pay-1", 9),
        issue(None, None, sheet="Пакет"),
        issue("f-main", 4, sheet="Контрагенты"),
        issue("unknown-file", 2),
    )
    rows = data_rows(open_report(AnalysisReport(meta, (), issues))["Качество данных"])
    assert [row[2] for row in rows] == [
        "лист «Пакет»",
        "лист «Платежи», строка 2",  # an id the report does not know stays unnamed
        "файл «Контрагенты», лист «Контрагенты», строка 4",
        "файл «Платежи за 01.06.2026–30.06.2026, файл 1», лист «Платежи», строка 9",
        "файл «Платежи за 01.07.2026–31.08.2026, файл 2», лист «Платежи», строка 3",
    ]


def test_about_sheet_names_run_mode_rules_and_limits():
    rows = dict(data_rows(open_report(live_report([]))["О проверке"])[:10])
    assert rows["ID проверки"] == "run-1"
    assert rows["Режим"] == "Реальные данные"
    assert rows["Оценки приоритета"].startswith("по правилам версии")
    assert rows["Внешние данные не старее"] == "18.09.2026 12:00 UTC"
    assert rows["Версия ИИ"] == "ИИ-пояснения не используются"


def test_about_sheet_names_what_is_not_checked_automatically():
    """S7-03: account blocks are a manual check, said with the page to open."""
    rows = data_rows(open_report(live_report([]))["О проверке"])
    manual = [value for key, value, *_ in rows if key == "Не проверяется автоматически"]
    assert len(manual) == 1
    assert manual[0].startswith("Блокировки счетов ФНС: не проверяются автоматически")
    assert "https://service.nalog.ru/bi.do" in manual[0]
    assert "БИК любого банка" in manual[0]


@pytest.mark.parametrize(
    "name", ['=HYPERLINK("http://example.org","x")', "=1+1", "+7 (900)", "@SUM(A1)"]
)
def test_text_that_looks_like_a_formula_stays_text(name):
    row = live_row(INN_B, Decimal("1.00"), 1, name=name)
    cell = open_report(live_report([row]))["Приоритеты"]["B3"]
    assert cell.data_type == "s"
    assert cell.value == name


def test_empty_report_is_still_a_valid_workbook():
    workbook = open_report(live_report([]))
    assert data_rows(workbook["Приоритеты"]) == []
    assert data_rows(workbook["Качество данных"])[0][0] == "Замечаний нет"


# --- demo report from the S2-06 control package ---


def test_demo_report_is_marked_on_every_sheet_and_keeps_predefined_scores():
    workbook = open_report(demo_report(NOW))
    for name in SHEETS:
        assert "ДЕМО" in workbook[name]["A1"].value
        assert "оценки приоритета заданы заранее" in workbook[name]["A1"].value
    rows = data_rows(workbook["Приоритеты"])
    assert len(rows) == len(DEMO_SCORES) == 38
    expected = {score.inn: PRIORITY_LABELS[Priority(score.priority.value)] for score in DEMO_SCORES}
    assert {row[0]: row[4] for row in rows} == expected
    assert all(row[6] == DEMO_SCORE_NOTE for row in rows)
    about = dict(data_rows(workbook["О проверке"]))
    assert about["Режим"] == "ДЕМО — синтетические данные"
    assert about["Оценки приоритета"] == "демонстрационные, заданы заранее"


def test_demo_report_quality_sheet_has_the_control_package_issues():
    rows = data_rows(open_report(demo_report(NOW))["Качество данных"])
    categories = [row[0] for row in rows]
    assert categories.count("Ошибка строки") == 11
    assert categories.count("Предупреждение") == 1
    # The demo note is shown once in the title, not repeated 38 times as a data gap.
    assert all(row[3] != DEMO_SCORE_NOTE for row in rows)


# --- model validation ---


def test_demo_scores_are_rejected_outside_demo_mode():
    with pytest.raises(ValueError):
        ReportMeta("r", DAY, DataMode.LIVE, NOW, demo_scores=True)


def test_timestamps_need_a_timezone():
    with pytest.raises(ValueError):
        ReportMeta("r", DAY, DataMode.DEMO, datetime(2026, 9, 19))


def test_row_and_report_reject_mixed_or_duplicate_inns():
    line = CounterpartyRow(inn=INN_A)
    with pytest.raises(ValueError):
        ReportRow(line, demo_assessment(INN_B, Priority.LOW))
    row = ReportRow(line, demo_assessment(INN_A, Priority.LOW))
    meta = ReportMeta("r", DAY, DataMode.DEMO, NOW)
    with pytest.raises(ValueError):
        AnalysisReport(meta, (row, row))


def test_internal_signal_reasons_cite_ids_present_in_grounds():
    workbook = open_report(live_report([live_row(INN_B, Decimal("500.00"), 90)]))
    (priority,) = data_rows(workbook["Приоритеты"])
    assert priority[4] == PRIORITY_LABELS[Priority.HIGH]
    assert "internal-overdue" in priority[6]
    grounds = {row[0] for row in data_rows(workbook["Основания"])}
    assert "internal-overdue" in grounds


def test_about_sheet_says_when_no_source_was_queried():
    meta = ReportMeta("run-2", DAY, DataMode.DEMO, NOW)  # checked_at is None
    rows = dict(data_rows(open_report(AnalysisReport(meta))["О проверке"])[:10])
    assert rows["Внешние данные не старее"] == "внешние источники не запрашивались"
