"""S8-08: the monitoring list — INNs from a message or a file, up to 100, added and removed.

Agreed with the customer on 16.10.2026. The files follow the shapes he sent — the 1C
report with an «ИНН» column («образец»), the same report without one, and the report
with its header cut off («тпз 3»); every name and INN here is synthetic.
"""

import io

import pytest
from openpyxl import Workbook
from openpyxl.styles import Alignment

from claims_assistant.application.watchlist_files import (
    NO_INN_COLUMN,
    REPORT_WITHOUT_INN,
    nothing_overdue,
    read_watch_file,
)
from claims_assistant.domain.inn import InnProblem
from claims_assistant.domain.sheet_rules import CellError, parse_inn
from claims_assistant.domain.watchlist import (
    WATCH_LIMIT,
    WATCH_OVERDUE_DAYS,
    Candidate,
    Invalid,
    OverdueCut,
    add,
    parse_inns,
    remove,
)
from claims_assistant.infrastructure.excel.reader import OpenpyxlSheetReader
from claims_assistant.presentation.telegram.watchlist import format_added, format_removed

_WEIGHTS = (2, 4, 10, 3, 5, 9, 4, 6, 8)


def inn(n: int) -> str:
    """A valid synthetic company INN: nine digits from ``n`` and their control digit."""
    head = f"{77_0000_000 + n:09d}"
    control = sum(w * int(d) for w, d in zip(_WEIGHTS, head, strict=True)) % 11 % 10
    return head + str(control)


ONE, TWO, THREE = inn(1), inn(2), inn(3)
ENTREPRENEUR = "500100732259"  # a valid twelve-digit INN from the validator's own tests


def broken(value: str) -> str:
    """The same INN with its control digit wrong."""
    return value[:-1] + str((int(value[-1]) + 1) % 10)


# --- a message -------------------------------------------------------------------------


@pytest.mark.parametrize(
    "text",
    [
        f"{ONE} {TWO} {THREE}",
        f"{ONE},{TWO}; {THREE}",
        f"{ONE}\n{TWO}\n{THREE}",
        f"1. ООО «Ромашка» — {ONE}\n2) ИП Иванов, ИНН {TWO}\n3. {THREE} (долг 125 000 ₽)",
    ],
)
def test_every_inn_of_a_message_in_its_order(text):
    parsed = parse_inns(text)
    assert [c.inn for c in parsed.candidates] == [ONE, TWO, THREE]
    assert parsed.invalid == () and parsed.repeats == 0


def test_an_entrepreneur_counts_too():
    assert [c.inn for c in parse_inns(f"{ENTREPRENEUR} {ONE}").candidates] == [ENTREPRENEUR, ONE]


def test_dates_amounts_and_ogrn_are_not_read_as_inns():
    text = f"на 16.10.2026, сумма 1250000, ОГРН 1027700132195, {ONE}"
    assert [c.inn for c in parse_inns(text).candidates] == [ONE]


def test_a_wrong_inn_is_named_by_its_place_without_its_value():
    parsed = parse_inns(f"{ONE} {broken(TWO)} {THREE[:-1]} {THREE}")
    assert [c.inn for c in parsed.candidates] == [ONE, THREE]
    assert parsed.invalid == (
        Invalid(2, InnProblem.CHECKSUM),
        Invalid(3, InnProblem.NINE_DIGITS),  # nine digits: one lost, not skipped
    )
    text = format_added(add([], parsed.candidates), parsed.invalid)
    assert "№ 2: Контрольная цифра ИНН не совпадает" in text
    assert broken(TWO) not in text and THREE[:-1] + " " not in text


def test_repeats_are_counted_once():
    parsed = parse_inns(f"{ONE} {TWO} {ONE} {ONE}")
    assert [c.inn for c in parsed.candidates] == [ONE, TWO]
    assert parsed.repeats == 2


# --- the list --------------------------------------------------------------------------


def test_new_inns_are_added_to_the_list_and_old_ones_named():
    result = add([ONE], [Candidate(ONE), Candidate(TWO, "ООО «Два»")])
    assert result.added == (Candidate(TWO, "ООО «Два»"),)
    assert result.already == (Candidate(ONE),)
    assert result.total == 2


def test_the_list_stops_at_its_limit_and_names_what_did_not_fit():
    current = [inn(n) for n in range(10, 10 + WATCH_LIMIT - 1)]  # 99 on the list
    result = add(current, [Candidate(ONE), Candidate(TWO), Candidate(current[0])])
    assert result.added == (Candidate(ONE),)
    assert result.over_limit == (Candidate(TWO),)
    assert result.already == (Candidate(current[0]),)
    assert result.total == WATCH_LIMIT
    text = format_added(result)
    assert f"не больше {WATCH_LIMIT} организаций (1)" in text and TWO in text
    assert text.endswith(f"Сейчас на мониторинге: {WATCH_LIMIT} из {WATCH_LIMIT}.")


def test_nothing_on_the_list_is_dropped_to_make_room():
    full = [inn(n) for n in range(10, 10 + WATCH_LIMIT)]
    result = add(full, [Candidate(ONE)])
    assert result.added == () and result.over_limit == (Candidate(ONE),)
    assert result.total == WATCH_LIMIT


def test_removing_names_what_was_not_there():
    result = remove([ONE, TWO], [TWO, THREE, TWO])
    assert (result.removed, result.absent, result.total) == ((TWO,), (THREE,), 1)
    text = format_removed(result, limit=WATCH_LIMIT)
    assert text.splitlines()[0] == "Убрано с мониторинга: 1."
    assert THREE in text and text.endswith(f"Сейчас на мониторинге: 1 из {WATCH_LIMIT}.")


# --- a file ----------------------------------------------------------------------------


def workbook(rows, indents=None) -> bytes:
    book = Workbook()
    sheet = book.active
    for number, values in enumerate(rows):
        sheet.append(list(values))
        if indents is not None:
            sheet.cell(row=sheet.max_row, column=1).alignment = Alignment(indent=indents[number])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def sample_report(with_inn=True) -> bytes:
    """«образец» (16.10.2026): counterparty → contract, the INN in the column that holds
    «Срок действия до» on a contract's row."""
    overdue = "Сумма просроченной задолженности с вычетом н/р дней"
    legend = ["Контрагент", None, "ИНН" if with_inn else None, overdue, "Дней"]
    rows = [
        (0, ["Отчет по просроченным лизинговым платежам на дату"]),
        (0, legend),
        (0, ["Договор лизинга", None, "Срок действия до"]),
        (2, ["ООО «Ромашка»", None, ONE if with_inn else None, 900, 307]),
        (4, ["Дог. фин. лизинга № 1 от 01.02.2024", None, "01.02.2027", 600, 307]),
        (4, ["Дог. фин. лизинга № 2 от 01.03.2024", None, "01.03.2027", 300, 30]),
        (2, ["ИП Иванов", None, ENTREPRENEUR if with_inn else None, 250, 95]),
        (4, ["Дог. фин. лизинга № 3 от 01.04.2024", None, "01.04.2027", 250, 95]),
        (0, ["Итого", None, None, 1150]),
    ]
    return workbook([values for _i, values in rows], [indent for indent, _v in rows])


def read(data: bytes, **options):
    return read_watch_file(data, OpenpyxlSheetReader(), **options)


def test_the_1c_report_with_an_inn_column_gives_names_and_inns():
    result = read(sample_report())
    assert result.candidates == (
        Candidate(ONE, "ООО «Ромашка»"),
        Candidate(ENTREPRENEUR, "ИП Иванов"),
    )
    assert result.problem is None and result.invalid == ()


def test_the_report_without_inns_says_names_are_not_enough():
    result = read(sample_report(with_inn=False))
    assert result.candidates == ()
    assert result.without_inn == 2
    assert result.problem == REPORT_WITHOUT_INN


def test_a_table_with_an_inn_column_and_names():
    data = workbook(
        [
            ["Список на мониторинг"],
            [],
            ["№", "Наименование", "ИНН"],
            [1, "ООО «Один»", int(ONE)],  # a number cell, as Excel keeps it
            [2, "ООО «Два»", TWO],
            [3, "ООО «Сломанный»", broken(THREE)],
            [4, "ООО «Один» ещё раз", ONE],
        ]
    )
    result = read(data)
    assert result.candidates == (Candidate(ONE, "ООО «Один»"), Candidate(TWO, "ООО «Два»"))
    assert result.invalid == (Invalid(6, InnProblem.CHECKSUM),)  # the row of the sheet
    assert result.repeats == 1
    text = format_added(add([], result.candidates), result.invalid, from_file=True)
    assert "строка 6: Контрольная цифра" in text and broken(THREE) not in text


def test_a_bare_column_of_inns_without_a_header():
    result = read(workbook([[ONE], [TWO], [int(THREE)]]))
    assert [c.inn for c in result.candidates] == [ONE, TWO, THREE]


def test_the_report_with_its_header_cut_off_is_read_by_the_inns_it_holds():
    """«тпз 3»: the INN on a company's row, a date in that column on its contracts."""
    data = workbook(
        [
            ["ООО «Ромашка»", ONE, "Иванова", "комментарий", "16.10.2026", 900, 307],
            ["Дог. № 1 от 01.02.2024", "01.02.2027", None, None, "16.10.2026", 600, 307],
            ["ИП Иванов", ENTREPRENEUR, "Петрова", None, "16.10.2026", 250, 95],
            ["Дог. № 3 от 01.04.2024", "01.04.2027", None, None, "16.10.2026", 250, 95],
        ]
    )
    result = read(data)
    assert [c.inn for c in result.candidates] == [ONE, ENTREPRENEUR]
    assert result.invalid == ()


def test_a_file_with_no_inns_at_all_says_what_to_send():
    result = read(workbook([["Наименование", "Сумма"], ["ООО «Один»", 100]]))
    assert result.candidates == () and result.problem == NO_INN_COLUMN


def test_a_long_list_of_problems_is_cut():
    invalid = tuple(Invalid(n, InnProblem.CHECKSUM) for n in range(1, 9))
    text = format_added(add([], []), invalid)
    assert "…и ещё 3" in text and "№ 6" not in text


def test_a_number_cell_that_lost_its_leading_zero_is_named_so_not_a_kpp():
    """Review A on #100: nine digits in a number cell are rather an INN of the regions
    01–09 whose zero Excel dropped (eleven for an entrepreneur's), not a КПП."""
    data = workbook(
        [
            ["Наименование", "ИНН"],
            ["ООО «Без нуля»", 123456789],
            ["ИП Без нуля", 12345678901],
            ["ООО «Текстом»", "123456789"],
        ]
    )
    result = read(data)
    assert result.invalid == (
        Invalid(2, InnProblem.LOST_ZERO),
        Invalid(3, InnProblem.LOST_ZERO),
        Invalid(4, InnProblem.NINE_DIGITS),  # typed as text: the КПП hint stays
    )
    text = format_added(add([], result.candidates), result.invalid, from_file=True)
    assert "строка 2: Похоже, Excel убрал ноль в начале ИНН" in text
    assert "123456789" not in text


@pytest.mark.parametrize("cell", [123456789, 12345678901, 123456789.0])
def test_the_package_reader_says_the_same_for_a_number_cell(cell):
    with pytest.raises(CellError, match="Excel убрал ноль"):
        parse_inn(cell)


# --- the 1C report cut by its overdue (S8-14) ------------------------------------------
# The customer, 17.10.2026 (question 8): «Всех с просрочкой больше 60 дней».


def overdue_report(*companies) -> bytes:
    """A synthetic «образец»: (name, INN or None, the days of each contract) per company,
    a contract's days None when its «Дней» cell is blank. The company's row shows its
    worst days, as the report's subtotal does."""
    legend = ["Контрагент", None, "ИНН", "Сумма просроченной задолженности", "Дней"]
    rows = [
        (0, ["Отчет по просроченным лизинговым платежам на дату"]),
        (0, legend),
        (0, ["Договор лизинга", None, "Срок действия до"]),
    ]
    number = 0
    for name, company_inn, days in companies:
        known = [d for d in days if d is not None]
        rows.append((2, [name, None, company_inn, 100 * len(days), max(known, default=None)]))
        for contract_days in days:
            number += 1
            contract = f"Дог. фин. лизинга № {number} от 01.02.2024"
            rows.append((4, [contract, None, "01.02.2027", 100, contract_days]))
    rows.append((0, ["Итого", None, None, 100 * number]))
    return workbook([values for _i, values in rows], [indent for indent, _v in rows])


def test_the_threshold_is_the_customers_sixty_days():
    assert WATCH_OVERDUE_DAYS == 60


def test_from_the_report_only_those_overdue_more_than_sixty_days_are_taken():
    data = overdue_report(
        ("ООО «Шестьдесят один»", ONE, [61]),
        ("ООО «Шестьдесят»", TWO, [60]),
        ("ООО «Без дней»", THREE, [None]),
        ("ООО «Худший договор»", inn(4), [30, 90]),  # the worst contract decides
    )
    result = read(data)
    assert result.candidates == (
        Candidate(ONE, "ООО «Шестьдесят один»"),
        Candidate(inn(4), "ООО «Худший договор»"),
    )
    assert result.cut == OverdueCut(days=60, total=4, taken=2, not_overdue=1, no_days=1)
    assert result.problem is None and result.without_inn == 0


def test_a_company_printed_twice_is_one_judged_by_its_worst_contract():
    """The report lists a client under every sales point: merged by ``merge_repeats``
    first, so 30 days under one point and 90 under another is one debtor of 90 days."""
    data = overdue_report(
        ("ООО «Дважды»", ONE, [30]),
        ("ООО «Недолгая»", TWO, [10]),
        ("ООО «Дважды»", ONE, [90]),
    )
    result = read(data)
    assert result.candidates == (Candidate(ONE, "ООО «Дважды»"),)
    assert result.cut == OverdueCut(days=60, total=2, taken=1, not_overdue=1, no_days=0)
    assert result.repeats == 0


def test_two_companies_of_one_name_with_different_inns_are_not_merged():
    """Two «ИП Иванов …» under different sales points are two companies: each is judged by
    its own days, and the one of 10 days does not go on as a debtor of 90."""
    data = overdue_report(
        ("ООО «Тёзка»", ONE, [10]),
        ("ООО «Другая»", THREE, [70]),
        ("ООО «Тёзка»", TWO, [90]),
    )
    result = read(data)
    assert result.candidates == (Candidate(THREE, "ООО «Другая»"), Candidate(TWO, "ООО «Тёзка»"))
    assert result.cut == OverdueCut(days=60, total=3, taken=2, not_overdue=1, no_days=0)


def test_a_company_with_days_only_on_some_contracts_is_judged_by_them():
    result = read(overdue_report(("ООО «Часть дней»", ONE, [None, 75])))
    assert result.candidates == (Candidate(ONE, "ООО «Часть дней»"),)
    assert result.cut.no_days == 0


def test_a_taken_company_without_an_inn_is_counted_not_guessed():
    """Only the taken ones are counted: a company cut out by its days is not on the list
    whatever its INN, and is already counted among «не взяты»."""
    data = overdue_report(
        ("ООО «С ИНН»", ONE, [90]),
        ("ООО «Без ИНН»", None, [120]),
        ("ООО «Без ИНН, недолгая»", None, [10]),
    )
    result = read(data)
    assert result.candidates == (Candidate(ONE, "ООО «С ИНН»"),)
    assert result.without_inn == 1
    assert result.cut == OverdueCut(days=60, total=3, taken=2, not_overdue=1, no_days=0)


def test_nobody_overdue_enough_is_said_plainly():
    data = overdue_report(("ООО «Шестьдесят»", ONE, [60]), ("ООО «Без дней»", TWO, [None]))
    result = read(data)
    assert result.candidates == ()
    assert result.cut == OverdueCut(days=60, total=2, taken=0, not_overdue=1, no_days=1)
    assert result.problem == nothing_overdue(result.cut)
    assert result.problem == (
        "Ни у одной из 2 организаций отчёта нет просрочки больше 60 дн., а из отчёта на "
        "мониторинг берутся только такие — добавлять некого. Без дней просрочки в отчёте: "
        "1 организация. Если кто-то из них всё же нужен, пришлите его ИНН сообщением."
    )


@pytest.mark.parametrize(
    ("days", "why"),
    [(60, "нет просрочки больше 60 дн."), (None, "нет дней просрочки")],
)
def test_nobody_of_a_one_company_report_is_said_of_it_alone(days, why):
    result = read(overdue_report(("ООО «Одна»", ONE, [days])))
    assert result.candidates == ()
    assert result.problem == (
        f"У единственной организации отчёта {why}, а из отчёта на мониторинг берутся только "
        "должники с просрочкой больше 60 дн. — добавлять некого. Если она всё же нужна, "
        "пришлите её ИНН сообщением."
    )


def test_a_report_without_inns_says_so_before_any_cut():
    result = read(overdue_report(("ООО «Без ИНН»", None, [10]), ("ИП Без ИНН", None, [None])))
    assert result.problem == REPORT_WITHOUT_INN and result.without_inn == 2


@pytest.mark.parametrize(
    ("overdue_days", "taken"),
    [
        (30, [ONE, TWO, THREE]),
        (60, [ONE, TWO]),
        (89, [ONE, TWO]),
        (90, [ONE]),
        (0, [ONE, TWO, THREE]),
    ],
)
def test_the_threshold_from_the_settings_moves_the_cut(overdue_days, taken):
    data = overdue_report(
        ("ООО «Сто»", ONE, [100]), ("ООО «Девяносто»", TWO, [90]), ("ООО «Сорок»", THREE, [40])
    )
    result = read(data, overdue_days=overdue_days)
    assert [c.inn for c in result.candidates] == taken
    assert result.cut.days == overdue_days


def test_a_table_and_a_bare_column_are_taken_whole_whatever_their_days():
    """The user's own list is his choice: only the 1C report is cut."""
    table = workbook(
        [["Наименование", "ИНН", "Дней"], ["ООО «Один»", ONE, 10], ["ООО «Два»", TWO, None]]
    )
    for data in (table, workbook([[ONE], [TWO]])):
        result = read(data, overdue_days=60)
        assert [c.inn for c in result.candidates] == [ONE, TWO]
        assert result.cut is None and result.problem is None


def test_the_reply_says_how_the_report_was_cut():
    cut = OverdueCut(days=60, total=318, taken=121, not_overdue=190, no_days=7)
    current = [inn(n) for n in range(10, 10 + WATCH_LIMIT - 2)]  # 98 on the list
    result = add(current, [Candidate(ONE), Candidate(TWO), Candidate(THREE)])
    text = format_added(result, from_file=True, cut=cut, without_inn=1)
    assert text.splitlines()[:4] == [
        "Добавлено на мониторинг: 2.",
        "В отчёте 318 организаций, просрочка больше 60 дн. — у 121 организации.",
        "Не взяты — просрочка не больше 60 дн.: 190 организаций.",
        "Не взяты — в отчёте нет дней просрочки: 7 организаций.",
    ]
    assert "Без ИНН среди отобранных: 1 организация — по названию" in text
    assert "Без ИНН в отчёте" not in text
    assert f"не больше {WATCH_LIMIT} организаций (1)" in text
    assert text.endswith(f"Сейчас на мониторинге: {WATCH_LIMIT} из {WATCH_LIMIT}.")


def test_the_reply_names_only_the_reasons_there_are():
    cut = OverdueCut(days=90, total=3, taken=3, not_overdue=0, no_days=0)
    text = format_added(add([], [Candidate(ONE)]), cut=cut)
    assert "В отчёте 3 организации, просрочка больше 90 дн. — у 3 организаций." in text
    assert "Не взяты" not in text
    assert "В отчёте" not in format_added(add([], [Candidate(ONE)]))
    assert "Без ИНН в отчёте: 1 организация" in format_added(add([], []), without_inn=1)
