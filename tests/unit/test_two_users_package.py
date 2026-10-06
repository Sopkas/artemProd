"""S6-03: the files of the two-user manual run do what the scenarios say they do.

The manual run is two people and a live bot; what they are told to expect has to be true
before they start. So the files are given to the application the way the dialog gives
them — accepted into a check of each owner, run by the worker with the real pipeline — and
every expectation of ``docs/acceptance-scenarios.md`` is held here: the priorities, that a
report and a card hold the owner's data only, and that the large file is accepted whole.
"""

from datetime import timedelta
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from claims_assistant.application.analysis_pipeline import AnalysisPipeline
from claims_assistant.application.check_package import accept_counterparties, accept_ledger
from claims_assistant.application.external_guard import GuardedCompanyDataProvider, GuardPolicy
from claims_assistant.application.imports import import_counterparties
from claims_assistant.application.internal_context import internal_context
from claims_assistant.application.worker import RunWorker
from claims_assistant.domain.analysis import FileKind, RunStatus
from claims_assistant.domain.external import DataMode
from claims_assistant.domain.inn import validate_inn
from claims_assistant.infrastructure.cache.memory import TtlSnapshotCache
from claims_assistant.infrastructure.demo.company_data import DemoCompanyDataProvider
from claims_assistant.infrastructure.demo.two_users import (
    ANALYSIS_DATE,
    ANDREY,
    SERGEY,
    SHARED_INN,
    STOP_ROWS,
    USERS,
)
from claims_assistant.infrastructure.excel.reader import OpenpyxlSheetReader
from claims_assistant.infrastructure.excel.report import PRIORITY_LABELS, build_report
from claims_assistant.infrastructure.excel.two_users_package import (
    STOP_FILE,
    build_two_users_package,
    counterparties_file,
    file_notes,
    interactions_file,
)
from claims_assistant.infrastructure.memory.analysis import InMemoryAnalysisRepository
from claims_assistant.infrastructure.storage.local import LocalFileStorage

OWNERS = {ANDREY.key: 101, SERGEY.key: 202}


@pytest.fixture(scope="module")
def package() -> dict[str, bytes]:
    return build_two_users_package()


def _rows(data: bytes, sheet: str) -> list[tuple]:
    book = load_workbook(BytesIO(data), read_only=True)
    return [row for row in book[sheet].iter_rows(min_row=3, values_only=True) if row[0]]


@pytest.fixture
async def checked(tmp_path, package):
    """Both users' checks, finished: each accepted from the user's own files and run by the
    worker in the order they were queued, as two people pressing «Запустить проверку»."""
    repository = InMemoryAnalysisRepository()
    files = LocalFileStorage(tmp_path / "uploads")
    reader = OpenpyxlSheetReader()
    deps = dict(repository=repository, files=files, reader=reader)
    runs = {}
    for user in USERS:
        owner = OWNERS[user.key]
        draft = await accept_counterparties(
            owner, ANALYSIS_DATE, DataMode.DEMO, package[counterparties_file(user)], **deps
        )
        assert draft.issues == (), draft.issues
        added = await accept_ledger(
            owner,
            draft.run.id,
            FileKind.INTERACTIONS,
            package[interactions_file(user)],
            coverage=None,
            **deps,
        )
        assert added.issues == () and added.rows == len(user.interactions)
        await repository.transition(owner, draft.run.id, RunStatus.QUEUED)
        runs[user.key] = draft.run.id
    provider = GuardedCompanyDataProvider(
        DemoCompanyDataProvider(), GuardPolicy(), TtlSnapshotCache(3600), mode=DataMode.DEMO
    )
    pipeline = AnalysisPipeline(
        files, reader, provider, repository, mode=DataMode.DEMO, build_report=build_report
    )
    worker = RunWorker(repository, pipeline)
    assert await worker.process_one() and await worker.process_one()
    return repository, files, reader, runs


@pytest.mark.parametrize("user", USERS, ids=[user.key for user in USERS])
async def test_each_user_gets_the_report_the_scenario_promises(checked, user):
    repository, files, _, runs = checked
    owner = OWNERS[user.key]
    run = await repository.get_run(owner, runs[user.key])
    assert run.status is RunStatus.COMPLETED
    data = files.read((await repository.get_report(owner, run.id)).stored_path)
    rows = {row[0]: row for row in _rows(data, "Приоритеты")}
    # Only the owner's companies, under the owner's names and with the promised priorities.
    assert set(rows) == set(user.inns)
    for row in user.rows:
        assert row.name in rows[row.inn]
        assert PRIORITY_LABELS[user.expected[row.inn]] in rows[row.inn], row.inn
    # «Хронология» holds the owner's interactions and nobody else's.
    comments = [cell for row in _rows(data, "Хронология") for cell in row if isinstance(cell, str)]
    spoken = [text for text in comments if text.startswith("Запись ")]
    assert sorted(spoken) == sorted(item.comment for item in user.interactions)


async def test_the_shared_inn_shows_each_user_their_own_figures(checked):
    """The sharpest check of the manual run: one INN in both packages. What the card of it
    reads for a user is that user's row and that user's interaction."""
    repository, files, reader, _ = checked
    seen = {}
    for user in USERS:
        context = await internal_context(OWNERS[user.key], SHARED_INN, repository, files, reader)
        mine = next(row for row in user.rows if row.inn == SHARED_INN)
        assert context is not None and context.row.debt == mine.debt
        assert context.row.name == mine.name
        assert [item.comment for item in context.interactions] == [
            item.comment for item in user.interactions if item.inn == SHARED_INN
        ]
        seen[user.key] = context.row.debt
    assert seen[ANDREY.key] != seen[SERGEY.key]  # otherwise the INN would prove nothing


async def test_a_company_of_the_other_user_has_no_internal_data(checked):
    repository, files, reader, _ = checked
    for user, other in ((ANDREY, SERGEY), (SERGEY, ANDREY)):
        foreign = next(inn for inn in other.inns if inn != SHARED_INN)
        assert await internal_context(OWNERS[user.key], foreign, repository, files, reader) is None


def test_the_packages_differ_in_everything_but_the_shared_inn():
    assert set(ANDREY.inns) & set(SERGEY.inns) == {SHARED_INN}
    assert not {row.name for row in ANDREY.rows} & {row.name for row in SERGEY.rows}
    for user in USERS:
        assert set(user.expected) == set(user.inns)
        assert all(user.owner[:-1] in row.name for row in user.rows)  # «Андрея», «Сергея»
        assert all(
            item.comment.startswith(f"Запись {user.owner[:-1]}") for item in user.interactions
        )
        # A high, a medium and a low in each report: a swapped row would stand out.
        assert len(set(user.expected.values())) == 3


def test_nothing_in_the_files_can_be_a_real_organisation(package):
    """Valid by the checksum, so the product takes them, yet no tax office has the code
    0000: with a live source left on by mistake nothing real would be asked about."""
    reader = OpenpyxlSheetReader()
    for name, data in package.items():
        if "kontragenty" not in name:
            continue
        result = import_counterparties(reader, data, analysis_date=ANALYSIS_DATE)
        assert result.issues == (), (name, result.issues[:3])
        for row in result.rows:
            assert validate_inn(row.inn) == row.inn and row.inn.startswith("0000")


def test_the_stop_file_keeps_a_check_busy(package):
    """Ten thousand different companies, every one accepted for the analysis date: a demo
    check of a few rows ends before the bot can be stopped by hand."""
    result = import_counterparties(
        OpenpyxlSheetReader(), package[STOP_FILE], analysis_date=ANALYSIS_DATE
    )
    assert result.issues == ()
    assert len({row.inn for row in result.rows}) == STOP_ROWS == 10_000


def test_another_analysis_date_rejects_every_row(package):
    """«Сегодня» instead of the date typed by hand: every row carries the cut-off date."""
    wrong = import_counterparties(
        OpenpyxlSheetReader(),
        package[counterparties_file(ANDREY)],
        analysis_date=ANALYSIS_DATE + timedelta(days=5),
    )
    assert wrong.rows == () and len(wrong.issues) == len(ANDREY.rows)


def test_the_scenarios_name_the_files_the_script_builds(package):
    doc = Path(__file__).resolve().parents[2] / "docs" / "acceptance-scenarios.md"
    text = doc.read_text(encoding="utf-8")
    assert set(package) == set(file_notes())
    for name in package:
        assert f"`{name}`" in text, name
    assert f"**`{ANALYSIS_DATE:%d.%m.%Y}`**" in text
    assert f"`{SHARED_INN}`" in text
    for user in USERS:
        for row in user.rows:
            # The table of expected priorities has a row per company of each owner.
            line = next(
                (
                    line
                    for line in text.splitlines()
                    if line.startswith(f"| {user.owner} |") and f"`{row.inn}`" in line
                ),
                "",
            )
            assert PRIORITY_LABELS[user.expected[row.inn]] in line, (user.key, row.inn)
