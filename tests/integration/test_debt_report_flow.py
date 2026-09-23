"""S7-01 (presentation): the overdue report from the package to the report and the card.

The file is the customer's own 1C print, and it names no INN — the whole path here is
about the link by company name holding all the way: through the package, the saved import
step, the report's sheets and the card of one company. Values are synthetic.
"""

import io
from datetime import UTC, date, datetime
from decimal import Decimal

import pytest
from openpyxl import Workbook
from openpyxl.styles import Alignment

from claims_assistant.application.analysis_queue import RunOutcome
from claims_assistant.application.check_package import (
    accept_counterparties,
    accept_debt_report,
)
from claims_assistant.application.internal_context import internal_context
from claims_assistant.domain.analysis import FileKind, RunStatus
from claims_assistant.domain.external import DataMode
from claims_assistant.infrastructure.demo.company_data import DemoCompanyDataProvider
from claims_assistant.infrastructure.excel.counterparties import build_counterparties_template
from claims_assistant.infrastructure.excel.reader import OpenpyxlSheetReader
from claims_assistant.infrastructure.memory.analysis import InMemoryAnalysisRepository
from claims_assistant.infrastructure.storage.local import LocalFileStorage
from claims_assistant.presentation.telegram.card import format_card
from tests.unit.test_analysis_pipeline import OWNER, guard, pipeline, sheet_rows

DAY = date(2026, 9, 1)
NOW = datetime(2026, 9, 1, 12, 0, tzinfo=UTC)
INN_1, INN_2 = "7707083893", "7710140679"
# The demo provider calls every company the same, and two companies of one name are
# exactly the case where contracts must NOT be attributed (tests/unit/test_contract_link).
# So the source here names them apart, as a real one would.
NAMES = {INN_1: "ООО «Ромашка»", INN_2: "АО «Василёк»"}
DEMO_NAME = NAMES[INN_1]


class NamingProvider:
    """The demo source with a name of its own for each company."""

    def __init__(self):
        self._inner = DemoCompanyDataProvider()

    async def fetch(self, request):
        snapshots = await self._inner.fetch(request)
        return tuple(_renamed(snapshot, NAMES.get(request.inn)) for snapshot in snapshots)


def _renamed(snapshot, name):
    from dataclasses import replace

    from claims_assistant.domain.external import FactKind

    if name is None:
        return snapshot
    facts = tuple(
        replace(fact, value=name) if fact.kind is FactKind.COMPANY_NAME else fact
        for fact in snapshot.facts
    )
    return replace(snapshot, facts=facts)


def debt_report(counterparties=None) -> bytes:
    """The print as 1C lays it out: three levels in one column, told apart by indent."""
    counterparties = counterparties or [
        (DEMO_NAME, [("Дог-1", 307, 1_200_000.0, "Экскаватор"), ("Дог-2", 30, 250_000.0, "Кран")])
    ]
    book = Workbook()
    sheet = book.active
    rows: list[tuple[int, list]] = [
        (0, ["Отчет по просроченным лизинговым платежам на дату"]),
        (0, ["Точка продаж", None, "Предмет лизинга", "Сумма", "Дней"]),
        (0, ["Контрагент"]),
        (0, ["Договор лизинга"]),
        (0, ["Департамент регионального развития"]),
    ]
    for name, contracts in counterparties:
        rows.append((2, [name]))
        for contract, days, amount, subject in contracts:
            rows.append((4, [contract, None, subject, amount, days]))
    rows.append((0, ["Итого", None, None, 1_450_000.0, None]))
    for indent, values in rows:
        sheet.append(values)
        sheet.cell(row=sheet.max_row, column=1).alignment = Alignment(indent=indent)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def deps(tmp_path):
    return {
        "repository": InMemoryAnalysisRepository(),
        "files": LocalFileStorage(tmp_path / "uploads"),
        "reader": OpenpyxlSheetReader(),
    }


async def package(deps, report=None):
    accepted = await accept_counterparties(
        OWNER,
        DAY,
        DataMode.DEMO,
        build_counterparties_template(),
        repository=deps["repository"],
        files=deps["files"],
        reader=deps["reader"],
    )
    result = await accept_debt_report(
        OWNER,
        accepted.run.id,
        report if report is not None else debt_report(),
        repository=deps["repository"],
        files=deps["files"],
        reader=deps["reader"],
    )
    return result


async def finish(deps, run):
    """The worker is what marks a run finished; the card only looks at finished ones."""
    await deps["repository"].transition(OWNER, run.id, RunStatus.QUEUED)
    await deps["repository"].claim_next()
    await deps["repository"].finish(run.id, RunOutcome(RunStatus.COMPLETED))


async def run_check(deps):
    run = (await package(deps)).run
    outcome = await pipeline(deps["files"], deps["repository"], guard(NamingProvider())).process(
        run
    )
    artifact = await deps["repository"].get_report(OWNER, run.id)
    return outcome, deps["files"].read(artifact.stored_path)


async def test_the_report_is_attached_to_the_package_as_its_own_kind(deps):
    result = await package(deps)
    assert result.rows == 1  # one counterparty came out of the print
    run = await deps["repository"].get_run(OWNER, result.run.id)
    kinds = [file.kind for file in run.files]
    assert kinds == [FileKind.COUNTERPARTIES, FileKind.DEBT_REPORT]


async def test_the_contracts_reach_the_report_and_the_priorities_sheet(deps):
    _, data = await run_check(deps)
    contracts = [row for row in sheet_rows(data, "Договоры") if row[0]]
    assert [(row[0], row[2], row[3]) for row in contracts] == [
        (INN_1, "Дог-1", 307),  # worst overdue first
        (INN_1, "Дог-2", 30),
    ]
    assert contracts[0][4] == Decimal("1200000.00") and contracts[0][5] == "Экскаватор"
    priorities = {row[0]: row for row in sheet_rows(data, "Приоритеты")}
    assert priorities[INN_1][10] == 2 and priorities[INN_1][11] == 307
    # The other company has no contracts in the report: empty, not zero — «не знаем».
    assert priorities[INN_2][10] is None and priorities[INN_2][11] is None


async def test_a_name_that_matches_nobody_is_named_in_the_quality_sheet(deps):
    report = debt_report([("ООО «Совсем другая»", [("Дог-9", 10, 1.0, "Станок")])])
    run = (await package(deps, report)).run
    await pipeline(deps["files"], deps["repository"], guard(NamingProvider())).process(run)
    artifact = await deps["repository"].get_report(OWNER, run.id)
    data = deps["files"].read(artifact.stored_path)
    quality = " ".join(str(c) for row in sheet_rows(data, "Качество данных") for c in row if c)
    assert "не отнесены ни к одной организации" in quality
    contracts = sheet_rows(data, "Договоры")
    assert contracts == [("Договоры не приложены", None, None, None, None, None, None, None)]


async def test_a_resumed_check_keeps_the_contracts_without_reading_the_file_again(deps):
    """The import step carries them, like every other file of the package (S3-03)."""
    run = (await package(deps)).run
    await pipeline(deps["files"], deps["repository"], guard(NamingProvider())).process(run)

    class Blind:
        def read(self, path):
            raise AssertionError("the package must not be read again")

        def save(self, run_id, data):
            return deps["files"].save(run_id, data)

        def remove(self, path):
            deps["files"].remove(path)

    blind = Blind()
    outcome = await pipeline(blind, deps["repository"], guard(NamingProvider())).process(
        await deps["repository"].get_run(OWNER, run.id)
    )
    assert outcome.status.value in ("completed", "partial")


async def test_the_card_shows_the_contracts_of_this_company(deps):
    from claims_assistant.application.check_company import check_company

    run = (await package(deps)).run
    await pipeline(deps["files"], deps["repository"], guard(NamingProvider())).process(run)
    await finish(deps, run)
    check = await check_company(INN_1, guard(NamingProvider()))
    context = await internal_context(
        OWNER,
        INN_1,
        deps["repository"],
        deps["files"],
        deps["reader"],
        company_name=DEMO_NAME,
    )
    assert [contract.name for contract in context.contracts] == ["Дог-1", "Дог-2"]
    card = format_card(check, context)
    assert "Договоры с просрочкой: 2" in card
    assert "Дог-1 · 307 дн." in card and "Экскаватор" in card


async def test_without_the_company_name_the_card_shows_no_contracts(deps):
    """Nothing to match against is «не знаем», not «привяжем как получится»."""
    run = (await package(deps)).run
    await pipeline(deps["files"], deps["repository"], guard(NamingProvider())).process(run)
    await finish(deps, run)
    context = await internal_context(
        OWNER, INN_1, deps["repository"], deps["files"], deps["reader"], company_name=None
    )
    assert context.contracts == ()
