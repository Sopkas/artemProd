"""S6-03: two users on one real stack — SQLite, local files, the dispatcher, the worker
with the full pipeline, the notifier. Nothing is mocked but Telegram's transport and the
external source (demo provider) and the model (stub).

Scenarios, each repeatable on its own:
1. two users go the whole way and each gets only their own report;
2. a user cannot reach another user's report or status, on any path;
3. the process dies in the middle of an analysis and a fresh process finishes it
   without repeating the import, the external requests or the model;
4. a report is delivered again on request, and a failed delivery is retried.
"""

import asyncio
from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from aiogram.methods import SendDocument, SendMessage
from openpyxl import load_workbook

from claims_assistant.application.ai_guard import GuardedRecommendationProvider
from claims_assistant.application.analysis_pipeline import AnalysisPipeline
from claims_assistant.application.analysis_repository import RunNotFound
from claims_assistant.application.external_guard import GuardedCompanyDataProvider, GuardPolicy
from claims_assistant.application.report_delivery import fetch_report
from claims_assistant.application.worker import RunWorker
from claims_assistant.domain.analysis import RunStatus
from claims_assistant.domain.counterparties import CounterpartyRow
from claims_assistant.domain.external import DataMode
from claims_assistant.domain.steps import DeliveryStatus
from claims_assistant.infrastructure.cache.memory import TtlSnapshotCache
from claims_assistant.infrastructure.demo.company_data import DemoCompanyDataProvider
from claims_assistant.infrastructure.excel.counterparties import build_counterparties_template
from claims_assistant.infrastructure.excel.reader import OpenpyxlSheetReader
from claims_assistant.infrastructure.excel.report import build_report
from claims_assistant.infrastructure.llm.stub import StubRecommendationProvider
from claims_assistant.infrastructure.persistence.sqlite import open_sqlite_repository
from claims_assistant.infrastructure.storage.local import LocalFileStorage
from claims_assistant.presentation.telegram import texts
from claims_assistant.presentation.telegram.handlers import create_dispatcher
from claims_assistant.presentation.telegram.notifier import TelegramRunNotifier

ANNA, BORIS = 42, 43  # two allowed users
STRANGER = 99
DAY = date(2026, 9, 1)
ANNA_ROWS = (
    CounterpartyRow(inn="7707083893", name="ООО Анны", debt=Decimal("1000.00"), cutoff_date=DAY),
    CounterpartyRow(inn="7710140679", debt=Decimal("5.00"), cutoff_date=DAY, overdue_days=90),
)
BORIS_ROWS = (
    CounterpartyRow(inn="1234567894", name="ООО Бориса", debt=Decimal("77.00"), cutoff_date=DAY),
)


class Stack:
    """Everything the service runs with, on a temporary directory."""

    def __init__(self, root, bot, *, external=None, explainer=None) -> None:
        self.root = root
        self.bot = bot
        self.repository = open_sqlite_repository(root / "claims.sqlite3")
        self.files = LocalFileStorage(root / "uploads")
        self.reader = OpenpyxlSheetReader()
        self.external = external or DemoCompanyDataProvider()
        self.calls: list[str] = []
        self.provider = GuardedCompanyDataProvider(
            _Counting(self.external, self.calls),
            GuardPolicy(),
            TtlSnapshotCache(0),  # no cache: every fetch is visible in `calls`
            mode=DataMode.DEMO,
        )
        self.model = explainer or StubRecommendationProvider()
        self.dispatcher = create_dispatcher(
            frozenset({ANNA, BORIS}),
            self.provider,
            repository=self.repository,
            files=self.files,
            reader=self.reader,
            mode=DataMode.DEMO,
        )
        self.pipeline = AnalysisPipeline(
            self.files,
            self.reader,
            self.provider,
            self.repository,
            mode=DataMode.DEMO,
            build_report=build_report,
            explainer=GuardedRecommendationProvider(self.model),
        )
        self.notifier = TelegramRunNotifier(bot, self.repository, self.files)
        self.worker = RunWorker(self.repository, self.pipeline, notifier=self.notifier)

    def close(self) -> None:
        self.repository.close()


async def asyncy_gather(task) -> None:
    await asyncio.gather(task, return_exceptions=True)


class _Counting:
    def __init__(self, inner, calls: list[str]) -> None:
        self.inner = inner
        self.calls = calls

    async def fetch(self, request):
        self.calls.append(request.inn)
        return await self.inner.fetch(request)


@pytest.fixture
def stack(tmp_path, bot):
    built = Stack(tmp_path, bot)
    yield built
    built.close()


# --- Telegram helpers --------------------------------------------------------------


def _document(name: str, data: bytes) -> dict:
    return {"file_id": name, "file_unique_id": name, "file_name": name, "file_size": len(data)}


async def say(stack: Stack, update_factory, user: int, text: str) -> SendMessage:
    before = len(stack.bot.session.calls)
    await stack.dispatcher.feed_update(stack.bot, update_factory(text, user_id=user))
    replies = [c for c in stack.bot.session.calls[before:] if isinstance(c, SendMessage)]
    assert replies, "the bot answered nothing"
    return replies[-1]


async def upload(stack: Stack, update_factory, user: int, name: str, data: bytes) -> SendMessage:
    stack.bot.session.files[name] = data
    before = len(stack.bot.session.calls)
    await stack.dispatcher.feed_update(
        stack.bot, update_factory(None, user_id=user, document=_document(name, data))
    )
    replies = [c for c in stack.bot.session.calls[before:] if isinstance(c, SendMessage)]
    return replies[-1]


async def launch_check(stack: Stack, update_factory, user: int, rows) -> str:
    """The user's whole dialog up to «Запустить проверку»; returns the run id."""
    await say(stack, update_factory, user, "Новая проверка")
    await say(stack, update_factory, user, "01.09.2026")
    reply = await upload(
        stack, update_factory, user, f"list-{user}.xlsx", build_counterparties_template(rows)
    )
    assert texts.CHECK_SUMMARY_TITLE in reply.text
    reply = await say(stack, update_factory, user, "Запустить проверку")
    assert reply.text.startswith(texts.CHECK_QUEUED_PREFIX)
    runs = await stack.repository.list_runs(user)
    assert runs[0].status == RunStatus.QUEUED
    return runs[0].id


def sent_to(stack: Stack, user: int, kind=SendMessage) -> list:
    return [c for c in stack.bot.session.calls if isinstance(c, kind) and c.chat_id == user]


def report_inns(data: bytes) -> set[str]:
    workbook = load_workbook(BytesIO(data), read_only=True)
    rows = workbook["Приоритеты"].iter_rows(min_row=3, values_only=True)
    return {row[0] for row in rows if row[0]}


# --- 1. two users, the whole way --------------------------------------------------


async def test_two_users_each_get_their_own_report(stack, update_factory):
    anna_run = await launch_check(stack, update_factory, ANNA, ANNA_ROWS)
    boris_run = await launch_check(stack, update_factory, BORIS, BORIS_ROWS)

    # One worker, runs in queue order; each notification goes to its owner only.
    assert await stack.worker.process_one() and await stack.worker.process_one()
    assert not await stack.worker.process_one()
    anna_docs, boris_docs = sent_to(stack, ANNA, SendDocument), sent_to(stack, BORIS, SendDocument)
    # The last document of each user is the report (the earlier one is the template).
    assert report_inns(anna_docs[-1].document.data) == {"7707083893", "7710140679"}
    assert report_inns(boris_docs[-1].document.data) == {"1234567894"}
    assert "Проверка завершена" in sent_to(stack, ANNA)[-1].text
    assert "Проверка завершена" in sent_to(stack, BORIS)[-1].text
    assert stack.calls == ["7707083893", "7710140679", "1234567894"]

    # Each user's status and report are their own.
    assert "01.09.2026" in (await say(stack, update_factory, ANNA, "Статус")).text
    for user, run_id, inns in (
        (ANNA, anna_run, {"7707083893", "7710140679"}),
        (BORIS, boris_run, {"1234567894"}),
    ):
        before = len(stack.bot.session.calls)
        await stack.dispatcher.feed_update(stack.bot, update_factory("/report", user_id=user))
        document = [c for c in stack.bot.session.calls[before:] if isinstance(c, SendDocument)][-1]
        assert document.chat_id == user and report_inns(document.document.data) == inns
        artifact = await stack.repository.get_report(user, run_id)
        assert artifact.delivery is DeliveryStatus.DELIVERED


# --- 2. no access to another user's data ------------------------------------------


async def test_a_user_never_sees_another_users_report_status_or_card_data(stack, update_factory):
    anna_run = await launch_check(stack, update_factory, ANNA, ANNA_ROWS)
    assert await stack.worker.process_one()

    # Boris has no checks: status and report say so, nothing of Anna's leaks.
    assert (await say(stack, update_factory, BORIS, "Статус")).text == texts.STATUS_EMPTY
    assert (await say(stack, update_factory, BORIS, "Отчёт")).text == texts.REPORT_EMPTY
    assert sent_to(stack, BORIS, SendDocument) == []
    # Nor through the application and repository, whatever the caller passes.
    assert await fetch_report(BORIS, stack.repository, stack.files) is None
    with pytest.raises(RunNotFound):
        await stack.repository.get_run(BORIS, anna_run)
    with pytest.raises(RunNotFound):
        await stack.repository.get_report(BORIS, anna_run)
    # The INN card shows Anna's internal data to Anna only.
    anna_card = await say(stack, update_factory, ANNA, "Проверить ИНН")
    anna_card = await say(stack, update_factory, ANNA, "7707083893")
    assert "Внутренние данные" in anna_card.text and "ООО Анны" not in anna_card.text
    await say(stack, update_factory, BORIS, "Проверить ИНН")
    boris_card = await say(stack, update_factory, BORIS, "7707083893")
    assert "Внутренние данные" not in boris_card.text
    # A stranger gets the denial and nothing else.
    reply = await say(stack, update_factory, STRANGER, "Отчёт")
    assert reply.text == texts.access_denied(STRANGER)


# --- 3. the process dies in the middle of an analysis ------------------------------


async def test_a_run_interrupted_mid_analysis_is_finished_by_the_next_process(
    tmp_path, bot, update_factory
):
    class Slow(DemoCompanyDataProvider):
        """Answers the first company, then hangs like a stuck network call."""

        def __init__(self) -> None:
            super().__init__()
            self.hung = asyncio.Event()

        async def fetch(self, request):
            if request.inn == "7710140679":
                self.hung.set()
                await asyncio.sleep(3600)
            return await super().fetch(request)

    slow = Slow()
    first = Stack(tmp_path, bot, external=slow)
    anna_run = await launch_check(first, update_factory, ANNA, ANNA_ROWS)
    task = asyncio.create_task(first.worker.run_forever())
    await asyncio.wait_for(slow.hung.wait(), timeout=5)
    # The process dies: the task is killed, nothing is finished, the run stays "running".
    task.cancel()
    await asyncy_gather(task)
    first.close()
    peek = open_sqlite_repository(tmp_path / "claims.sqlite3")
    try:
        assert (await peek.get_run(ANNA, anna_run)).status == RunStatus.RUNNING
    finally:
        peek.close()
    assert "Проверка завершена" not in " ".join(m.text for m in sent_to(first, ANNA))

    # A fresh process on the same data: recovery requeues, the run completes.
    second = Stack(tmp_path, bot)
    try:
        recovered = await second.worker.recover()
        assert [r.id for r in recovered] == [anna_run]
        assert await second.worker.process_one()
        run = await second.repository.get_run(ANNA, anna_run)
        assert run.status == RunStatus.COMPLETED and run.attempts == 2
        # The first company's section came from the saved step: only the second was fetched.
        assert second.calls == ["7710140679"]
        assert second.model.requests, "the model was asked in the second process"
        report = sent_to(second, ANNA, SendDocument)[-1]
        assert report_inns(report.document.data) == {"7707083893", "7710140679"}
        assert "Проверка завершена" in sent_to(second, ANNA)[-1].text
    finally:
        second.close()


# --- 4. delivery again, and after a failure ----------------------------------------


async def test_report_is_delivered_again_and_a_failed_delivery_is_retried(stack, update_factory):
    anna_run = await launch_check(stack, update_factory, ANNA, ANNA_ROWS)
    assert await stack.worker.process_one()
    first = sent_to(stack, ANNA, SendDocument)[-1].document.data

    # Again on request: the same file, no new analysis, no new model call.
    asked = len(stack.model.requests)
    for _ in range(2):
        before = len(stack.bot.session.calls)
        await stack.dispatcher.feed_update(stack.bot, update_factory("/report", user_id=ANNA))
        again = [c for c in stack.bot.session.calls[before:] if isinstance(c, SendDocument)][-1]
        assert again.document.data == first
    assert len(stack.model.requests) == asked and stack.calls == ["7707083893", "7710140679"]
    assert len(await stack.repository.list_runs(ANNA)) == 1

    # Telegram refuses the file once: the user is told to retry, the next try works.
    stack.bot.session.fail_once = RuntimeError("telegram is down")
    reply = await say(stack, update_factory, ANNA, "Отчёт")
    assert reply.text == texts.REPORT_SEND_FAILED
    assert (await stack.repository.get_report(ANNA, anna_run)).delivery is DeliveryStatus.FAILED
    before = len(stack.bot.session.calls)
    await stack.dispatcher.feed_update(stack.bot, update_factory("Отчёт", user_id=ANNA))
    assert [c for c in stack.bot.session.calls[before:] if isinstance(c, SendDocument)]
    assert (await stack.repository.get_report(ANNA, anna_run)).delivery is DeliveryStatus.DELIVERED


def test_scenarios_are_documented():
    """The manual counterpart of these scenarios lives in docs/acceptance-scenarios.md."""
    from pathlib import Path

    doc = Path(__file__).resolve().parents[2] / "docs" / "acceptance-scenarios.md"
    text = doc.read_text(encoding="utf-8")
    for anchor in ("два пользователя", "чужой отчёт", "остановка", "повторная доставка"):
        assert anchor in text.lower()
