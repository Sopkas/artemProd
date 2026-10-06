"""S6-04: the defence scenario of docs/demo-script.md, typed into the bot (review A on #61).

The unit tests of the package feed the pipeline directly and fill in the date and the period
themselves, so a step the presenter types by hand is never checked there. Here the dialog
is driven button by button with the files the script names, on the source the service builds
from the settings of the show, and the worker delivers the report as it does live. Only the
Telegram transport is replaced.
"""

from io import BytesIO

import pytest
from aiogram.methods import SendDocument
from openpyxl import load_workbook

from claims_assistant.application.analysis_pipeline import AnalysisPipeline
from claims_assistant.application.external_guard import GuardedCompanyDataProvider, GuardPolicy
from claims_assistant.application.worker import RunWorker
from claims_assistant.domain.external import DataMode
from claims_assistant.infrastructure.cache.memory import TtlSnapshotCache
from claims_assistant.infrastructure.demo.package import (
    ALARM_INN,
    INCOMPLETE_INN,
    ORDINARY_INN,
    WORSENING_INN,
)
from claims_assistant.infrastructure.excel.demo_package import (
    COUNTERPARTIES_FILE,
    HISTORY_FILE,
    INTERACTIONS_FILE,
    PAYMENTS_FILE,
    build_demo_package,
)
from claims_assistant.infrastructure.excel.reader import OpenpyxlSheetReader
from claims_assistant.infrastructure.excel.report import build_report
from claims_assistant.infrastructure.persistence.sqlite import open_sqlite_repository
from claims_assistant.infrastructure.storage.local import LocalFileStorage
from claims_assistant.presentation.telegram import texts
from claims_assistant.presentation.telegram.handlers import create_dispatcher
from claims_assistant.presentation.telegram.notifier import TelegramRunNotifier
from claims_assistant.runtime.app import company_source
from claims_assistant.runtime.settings import Settings

PRESENTER = 42
TOKEN = "123456789:synthetic_token_for_offline_tests_only"


def show_settings(tmp_path, monkeypatch, **overrides) -> Settings:
    """The .env of the show, as the script gives it."""
    values = {
        "TELEGRAM_BOT_TOKEN": TOKEN,
        "ALLOWED_TELEGRAM_IDS": str(PRESENTER),
        "DATA_PROVIDER": "demo",
        "DEMO_PACKAGE": "true",
        "AI_PROVIDER": "off",
    } | overrides
    for key, value in values.items():
        monkeypatch.setenv(key, value)
    return Settings.load(tmp_path / "missing.env")


class Show:
    """The bot as it runs at the defence: SQLite, files, worker, notifier."""

    def __init__(self, root, bot, settings: Settings) -> None:
        self.bot = bot
        self.repository = open_sqlite_repository(root / "claims.sqlite3")
        self.files = LocalFileStorage(root / "uploads")
        reader = OpenpyxlSheetReader()
        provider = GuardedCompanyDataProvider(
            company_source(settings), GuardPolicy(), TtlSnapshotCache(3600), mode=DataMode.DEMO
        )
        self.dispatcher = create_dispatcher(
            settings.allowed_ids,
            provider,
            repository=self.repository,
            files=self.files,
            reader=reader,
            mode=DataMode.DEMO,
        )
        pipeline = AnalysisPipeline(
            self.files,
            reader,
            provider,
            self.repository,
            mode=DataMode.DEMO,
            build_report=build_report,
        )
        notifier = TelegramRunNotifier(bot, self.repository, self.files)
        self.worker = RunWorker(self.repository, pipeline, notifier=notifier)
        self.package = build_demo_package()

    async def say(self, update_factory, text: str) -> str:
        await self.dispatcher.feed_update(self.bot, update_factory(text, user_id=PRESENTER))
        return self.bot.session.calls[-1].text

    async def send(self, update_factory, name: str) -> str:
        data = self.package[name]
        self.bot.session.files[name] = data
        document = {"file_id": name, "file_unique_id": name, "file_name": name}
        document["file_size"] = len(data)
        await self.dispatcher.feed_update(
            self.bot, update_factory(None, user_id=PRESENTER, document=document)
        )
        return self.bot.session.calls[-1].text

    async def run_the_script(self, update_factory, *, day="01.09.2026", period=None) -> None:
        """Steps 1 of the script: the date by hand, then every file of the package."""
        await self.say(update_factory, "Новая проверка")
        await self.say(update_factory, day)
        assert "организаций: 4" in await self.send(update_factory, COUNTERPARTIES_FILE)
        await self.say(update_factory, "Добавить платежи")
        await self.say(update_factory, "По нашему шаблону")
        await self.say(update_factory, period or "01.06.2026–01.09.2026")
        await self.send(update_factory, PAYMENTS_FILE)
        await self.say(update_factory, "Добавить историю долга")
        await self.send(update_factory, HISTORY_FILE)
        await self.say(update_factory, "Добавить взаимодействия")
        await self.send(update_factory, INTERACTIONS_FILE)
        assert (await self.say(update_factory, "Запустить проверку")).startswith(
            texts.CHECK_QUEUED_PREFIX
        )
        assert await self.worker.process_one()

    def report(self) -> dict[str, tuple]:
        """INN → its row on «Приоритеты» of the report the bot sent."""
        documents = [c for c in self.bot.session.calls if isinstance(c, SendDocument)]
        book = load_workbook(BytesIO(documents[-1].document.data), read_only=True)
        rows = book["Приоритеты"].iter_rows(min_row=3, values_only=True)
        return {row[0]: row for row in rows if row[0]}

    def close(self) -> None:
        self.repository.close()


@pytest.fixture
def show(tmp_path, bot, monkeypatch):
    built = Show(tmp_path, bot, show_settings(tmp_path, monkeypatch))
    yield built
    built.close()


async def test_the_script_gives_the_four_stories_it_promises(show, update_factory):
    await show.run_the_script(update_factory)
    rows = show.report()
    assert {inn: row[4] for inn, row in rows.items()} == {
        ORDINARY_INN: "Средний",
        ALARM_INN: "Высокий",
        WORSENING_INN: "Высокий",
        INCOMPLETE_INN: "Недостаточно данных",
    }
    # «Северный путь» is high because of the two optional files, not its own row.
    reasons = rows[WORSENING_INN][6]
    assert "internal-payments" in reasons and "internal-debt-history" in reasons
    # The end of the run is said as the script says it: partial, because of «Тихое».
    finished = [c.text for c in show.bot.session.calls if getattr(c, "text", None)][-1]
    assert "завершена частично" in finished
    assert "Организаций: 4, проверено полностью: 3" in finished


async def test_the_card_of_scene_two_shows_the_owners_own_data(show, update_factory):
    await show.run_the_script(update_factory)
    await show.say(update_factory, "Проверить ИНН")
    card = await show.say(update_factory, ORDINARY_INN)
    assert "Внутренние данные" in card and "Взаимодействия: 1" in card
    assert "ООО «Ромашка» (демо)" in card  # the package names the company, not «ДЕМО — …»


async def test_a_period_typed_one_day_short_is_what_the_script_warns_about(show, update_factory):
    """«Если что-то пойдёт не так»: the period must reach the analysis date."""
    await show.run_the_script(update_factory, period="01.06.2026–31.08.2026")
    assert show.report()[WORSENING_INN][4] == "Средний"


async def test_without_demo_package_the_alarm_is_only_medium(
    tmp_path, bot, monkeypatch, update_factory
):
    """A forgotten DEMO_PACKAGE shows on the screen, as the script says."""
    settings = show_settings(tmp_path, monkeypatch, DEMO_PACKAGE="false")
    plain = Show(tmp_path, bot, settings)
    try:
        await plain.run_the_script(update_factory)
        assert plain.report()[ALARM_INN][4] == "Средний"
    finally:
        plain.close()
