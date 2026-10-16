"""S8-09 (file, part B): the monitoring list as an Excel file — the «Мониторинг» button.

The chat shows how many are watched and the latest few; the whole list — up to a hundred
companies, with when each was added, last checked and last changed — is a file the
customer can sort, filter and keep. The INN column is text: an entrepreneur's twelve
digits and a leading zero survive the round trip into his own lists.
"""

import io
from collections.abc import Sequence
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from claims_assistant.domain.watchlist import WATCH_LIMIT, WatchEntry

SHEET_NAME = "Мониторинг"
TITLES = (
    "№",
    "ИНН",
    "Название",
    "На мониторинге с",
    "Последняя проверка",
    "Последнее изменение",
    "Дата изменения",
)
WIDTHS = (5, 14, 45, 17, 19, 70, 16)
DATE_FORMAT = "DD.MM.YYYY"


def build_watchlist(entries: Sequence[WatchEntry], as_of: date, limit: int = WATCH_LIMIT) -> bytes:
    """The list, newest first as added; an empty list is a file with the header only."""
    book = Workbook()
    sheet = book.active
    sheet.title = SHEET_NAME
    sheet.append([f"Список на мониторинге на {as_of:%d.%m.%Y}: {len(entries)} из {limit}"])
    sheet["A1"].font = Font(bold=True)
    sheet.append(list(TITLES))
    for cell in sheet[2]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ordered = sorted(entries, key=lambda entry: (entry.added_on, entry.inn), reverse=True)
    for number, entry in enumerate(ordered, start=1):
        sheet.append(
            [
                number,
                entry.inn,
                entry.name or "",
                entry.added_on,
                entry.last_checked_on,
                entry.last_change or "",
                entry.last_change_on,
            ]
        )
        row = sheet.max_row
        sheet.cell(row=row, column=2).number_format = "@"
        for column in (4, 5, 7):
            sheet.cell(row=row, column=column).number_format = DATE_FORMAT
        sheet.cell(row=row, column=6).alignment = Alignment(wrap_text=True, vertical="top")
    for column, width in enumerate(WIDTHS, start=1):
        sheet.column_dimensions[sheet.cell(row=2, column=column).column_letter].width = width
    sheet.freeze_panes = "C3"
    sheet.auto_filter.ref = f"A2:G{max(sheet.max_row, 2)}"
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
