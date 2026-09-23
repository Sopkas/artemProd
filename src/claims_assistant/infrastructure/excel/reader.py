"""S2-04: sheet reader with archive-size and row guards; .xls as well since S4-05.

The unpacked size is checked before openpyxl parses the archive, so a decompression
bomb is rejected without allocating its expanded content.

The customer's own exports come out of 1C in the old .xls format, so that is read too,
through ``xlrd``. One difference cannot be papered over: **a formula inside an .xls is
invisible to us**. The file keeps the value Excel computed last, and the library hands it
over as if it had been typed; in an .xlsx a formula is refused, because a saved value may
be stale. Callers tell the formats apart with ``is_legacy_xls`` and warn the user.
"""

import io
import zipfile
from datetime import datetime, time

import xlrd
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from claims_assistant.application.imports import (
    CorruptWorkbook,
    Sheet,
    SheetMissing,
    WorkbookTooLarge,
)
from claims_assistant.domain.counterparties import Cell, FormulaCell, ImportLimits
from claims_assistant.domain.export_1c import MAX_HEADER_ROWS as MAX_EXPORT_HEADER_ROWS

# Compound File Binary header: every .xls — and any other OLE2 document — starts with it.
_OLE2_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def is_legacy_xls(source: bytes) -> bool:
    """Whether these bytes are an old-format .xls rather than an .xlsx archive."""
    return source[:8] == _OLE2_MAGIC


class OpenpyxlSheetReader:
    """Read a named sheet from an .xlsx given as bytes, a path or a binary stream."""

    def read(self, source: object, sheet: str, limits: ImportLimits) -> Sheet:
        data = self._as_bytes(source)
        if is_legacy_xls(data):
            rows = self._legacy_rows(data, limits, sheet)
            header = rows[0] if rows else ()
            return header, [(index + 2, cells) for index, cells in enumerate(rows[1:])]
        self._guard_unpacked_size(data, limits)
        try:
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
        except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError, ValueError):
            raise CorruptWorkbook() from None
        try:
            if sheet not in workbook.sheetnames:
                raise SheetMissing(sheet)
            # Materialize rows before closing: read-only rows are tied to the open workbook.
            return self._extract(workbook[sheet], limits)
        finally:
            workbook.close()

    @staticmethod
    def _as_bytes(source: object) -> bytes:
        if isinstance(source, bytes):
            return source
        if isinstance(source, bytearray):
            return bytes(source)
        if isinstance(source, str):
            with open(source, "rb") as handle:
                return handle.read()
        if hasattr(source, "read"):
            return source.read()
        raise CorruptWorkbook()

    @staticmethod
    def _guard_unpacked_size(data: bytes, limits: ImportLimits) -> None:
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                total = sum(info.file_size for info in archive.infolist())
        except zipfile.BadZipFile:
            raise CorruptWorkbook() from None
        if total > limits.max_unpacked_bytes:
            raise WorkbookTooLarge()

    def read_rows(
        self, source: object, limits: ImportLimits, sheet: str | None = None
    ) -> tuple[tuple[Cell, ...], ...]:
        """Every row of one sheet, header block included (S4-05).

        The customer's own export prints a parameters block above the table and names the
        sheet «Лист_1», so neither «the first row is the header» nor a known sheet name
        holds; ``sheet=None`` takes the workbook's first sheet.
        """
        data = self._as_bytes(source)
        if is_legacy_xls(data):
            return self._legacy_rows(data, limits, sheet)
        self._guard_unpacked_size(data, limits)
        try:
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
        except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError, ValueError):
            raise CorruptWorkbook() from None
        try:
            if sheet is not None and sheet not in workbook.sheetnames:
                raise SheetMissing(sheet)
            worksheet = workbook[sheet] if sheet is not None else workbook[workbook.sheetnames[0]]
            rows = []
            for row in worksheet.iter_rows():
                rows.append(tuple(OpenpyxlSheetReader._cell(cell) for cell in row))
                if len(rows) > limits.max_rows + MAX_EXPORT_HEADER_ROWS:
                    break
            return tuple(rows)
        finally:
            workbook.close()

    @staticmethod
    def _extract(worksheet: object, limits: ImportLimits) -> Sheet:
        header: tuple[Cell, ...] = ()
        rows: list[tuple[int, tuple[Cell, ...]]] = []
        for offset, row in enumerate(worksheet.iter_rows()):
            cells = tuple(OpenpyxlSheetReader._cell(cell) for cell in row)
            if offset == 0:
                header = cells
                continue
            rows.append((offset + 1, cells))
            # One extra row past the limit lets the parser report the overflow.
            if len(rows) > limits.max_rows:
                break
        return header, rows

    @staticmethod
    def _legacy_rows(
        data: bytes, limits: ImportLimits, sheet: str | None
    ) -> tuple[tuple[Cell, ...], ...]:
        """Rows of an old-format .xls; the sheet by name, or the first one."""
        if len(data) > limits.max_unpacked_bytes:
            raise WorkbookTooLarge()
        try:
            book = xlrd.open_workbook(file_contents=data, on_demand=True)
        except Exception:  # xlrd raises its own family for every kind of broken file
            raise CorruptWorkbook() from None
        try:
            if sheet is not None and sheet not in book.sheet_names():
                raise SheetMissing(sheet)
            worksheet = book.sheet_by_name(sheet) if sheet is not None else book.sheet_by_index(0)
            rows = []
            for index in range(worksheet.nrows):
                rows.append(
                    tuple(
                        OpenpyxlSheetReader._legacy_cell(
                            worksheet.cell(index, column), book.datemode
                        )
                        for column in range(worksheet.ncols)
                    )
                )
                if len(rows) > limits.max_rows + MAX_EXPORT_HEADER_ROWS:
                    break
            return tuple(rows)
        finally:
            book.release_resources()

    @staticmethod
    def _legacy_cell(cell: object, datemode: int) -> Cell:
        """One .xls cell in the same types an .xlsx gives: str, int, float, date, bool.

        A formula is not visible here — xlrd reports the value the file carries — so an
        .xls cannot be checked for formulas the way an .xlsx is.
        """
        kind, value = cell.ctype, cell.value
        if kind == xlrd.XL_CELL_EMPTY or kind == xlrd.XL_CELL_BLANK:
            return None
        if kind == xlrd.XL_CELL_TEXT:
            # A text that starts with "=" is still refused, as in an .xlsx.
            return FormulaCell() if value.startswith("=") else value
        if kind == xlrd.XL_CELL_BOOLEAN:
            return bool(value)
        if kind == xlrd.XL_CELL_ERROR:
            # #REF!, #DIV/0! and the like carry no value. They come back as an empty cell,
            # so the parser reports the field as missing — with its row and column — which
            # is what a person needs to fix the file (A on #51).
            return None
        if kind == xlrd.XL_CELL_DATE:
            parts = xlrd.xldate_as_tuple(value, datemode)
            if parts[:3] == (0, 0, 0):
                return time(*parts[3:])
            moment = datetime(*parts)
            return moment.date() if parts[3:] == (0, 0, 0) else moment
        # Numbers come back as floats; whole ones become int, as openpyxl would give.
        return int(value) if float(value).is_integer() else value

    @staticmethod
    def _cell(cell: object) -> Cell:
        value = getattr(cell, "value", None)
        # Detect formulas by data type and, defensively, by a leading "=".
        if getattr(cell, "data_type", None) == "f" or (
            isinstance(value, str) and value.startswith("=")
        ):
            return FormulaCell()
        return value
